from io import BytesIO
from pptx import Presentation
from pptx.util import Inches, Pt

def build_presentation(req):
    prs=Presentation(); prs.core_properties.title=req.title; prs.core_properties.subject=req.topic
    slide=prs.slides.add_slide(prs.slide_layouts[0]); slide.shapes.title.text=req.title; slide.placeholders[1].text=f"Topic: {req.topic} | Audience: {req.audience}\nEducational study aid • verify facts with authoritative sources"
    sections=[("Learning objectives",[f"Define {req.topic}","Identify core concepts and terminology","Explain relevance to healthcare learning"]),("Core concepts",[f"Introduce the foundational principles of {req.topic}","Separate established evidence from assumptions","Use trusted sources for details"]),("Practical application",["Work through a learning example or case discussion","Consider context, limitations, and safety","Ask what evidence would support a conclusion"]),("Review and discussion",["Summarize three key takeaways","List unanswered questions","Plan further reading from authoritative sources"])]
    for i in range(req.slide_count-2):
        title,bullets=sections[i%len(sections)]; s=prs.slides.add_slide(prs.slide_layouts[1]);s.shapes.title.text=title
        tf=s.placeholders[1].text_frame;tf.clear()
        for j,b in enumerate(bullets):
            p=tf.paragraphs[0] if j==0 else tf.add_paragraph();p.text=b;p.level=0;p.font.size=Pt(22)
    s=prs.slides.add_slide(prs.slide_layouts[1]);s.shapes.title.text="References and responsible use";tf=s.placeholders[1].text_frame;tf.clear()
    for i,line in enumerate(["No specific papers or citations were supplied or fabricated.","Consult current official guidance, textbooks, and peer-reviewed literature.","This presentation is an educational aid, not medical advice."]):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=line;p.font.size=Pt(18)
    out=BytesIO();prs.save(out);out.seek(0);return out.getvalue()
